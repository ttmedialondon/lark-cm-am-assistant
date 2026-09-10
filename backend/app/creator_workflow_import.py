"""Bounded Excel preflight and idempotent, per-row intake; original files stay untouched."""

from __future__ import annotations

import hashlib
import io
import json
import re
import zipfile
from pathlib import PurePath
from uuid import UUID, uuid4

from openpyxl import load_workbook

from . import creator_workflow as workflow
from .db import connection

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ROWS = 3000
HEADERS = {
    "handle": {"creatorhandle", "handle", "达人handle", "tiktokhandle", "达人账号"},
    "cm": {"cm", "负责人"},
    "tap_cm": {"tapcm"},
    "open_id": {"larkopenid", "openid", "lark用户id"},
    "contact_id": {"contactid", "联系人id"},
    "relationship": {"relationship", "用户关系", "本人或代表"},
}


def _text(value):
    if value is None:
        return None
    text = str(value).strip()
    return None if text in {"", "-", "—", "–"} else text[:500]


def parse_workbook(content: bytes, filename: str) -> list[dict]:
    if not content or len(content) > MAX_FILE_BYTES or not filename.lower().endswith(".xlsx"):
        raise workflow.WorkflowInputError("请上传不超过 10 MB 的 .xlsx 文件。")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if (
                len(entries) > 2000
                or sum(e.file_size for e in entries) > 50 * 1024 * 1024
                or any(e.flag_bits & 1 for e in entries)
            ):
                raise ValueError("oversized workbook")
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            rows = []
            found_header = False
            # Lark-exported workbooks sometimes incorrectly declare an A1-only dimension.
            for sheet in book:
                sheet.reset_dimensions()
                columns = None
                for number, cells in enumerate(sheet.iter_rows(max_col=80), 1):
                    values = [_text(cell.value) for cell in cells]
                    if columns is None:
                        if number > 10:
                            break
                        labels = [re.sub(r"[\s_\-]", "", v or "").casefold() for v in values]
                        found = {
                            key: [i for i, name in enumerate(labels) if name in aliases]
                            for key, aliases in HEADERS.items()
                        }
                        if not found["handle"]:
                            continue
                        if any(len(indices) > 1 for indices in found.values()):
                            raise workflow.WorkflowInputError("表头有重复的 handle / CM / 用户 ID 列，请先区分。")
                        columns = {key: indices[0] for key, indices in found.items() if indices}
                        found_header = True
                        continue
                    if not any(values):
                        continue
                    if len(rows) >= MAX_ROWS:
                        raise workflow.WorkflowInputError("单次最多导入 3000 行，请拆分文件。")
                    row = {key: values[index] for key, index in columns.items()}
                    formula = any(cells[index].data_type == "f" for index in columns.values())
                    row.update(sheet=sheet.title, row=number, formula=formula)
                    rows.append(row)
            if rows:
                return rows
            if found_header:
                raise workflow.WorkflowInputError("已找到 handle 表头，但没有可导入的数据行。")
        finally:
            book.close()
    except workflow.WorkflowInputError:
        raise
    except Exception as exc:
        raise workflow.WorkflowInputError("无法读取 Excel，请另存为标准 .xlsx 后重试。") from exc
    raise workflow.WorkflowInputError("找不到 handle 表头。支持 creator handle、handle、达人 handle。")


def _aliases(conn):
    rows = conn.execute("""SELECT auth.id,auth.display_name,identity.name,identity.en_name
        FROM authorizations auth LEFT JOIN lark_identities identity ON identity.open_id=auth.open_id
        WHERE auth.status='active' AND auth.assignment_eligible""").fetchall()
    names = {}
    for row in rows:
        for value in (row["display_name"], row["name"], row["en_name"]):
            if value and value.strip():
                names.setdefault(value.strip().casefold(), set()).add(str(row["id"]))
        names.setdefault(str(row["id"]), set()).add(str(row["id"]))
    for row in conn.execute("""SELECT alias,responsible_id FROM creator_workflow_cm_aliases alias
        JOIN authorizations auth ON auth.id=alias.responsible_id
        WHERE auth.status='active' AND auth.assignment_eligible""").fetchall():
        names.setdefault(row["alias"].strip().casefold(), set()).add(str(row["responsible_id"]))
    return names


def resolve_rows(conn, scope, rows):
    aliases = _aliases(conn)
    result = []
    for source in rows:
        reasons = []
        handle = None
        try:
            handle = workflow._handle(source.get("handle"))
        except workflow.WorkflowInputError:
            reasons.append("handle_uncertain")
        if not handle:
            reasons.append("missing_handle")
        cm_values = [source.get("cm"), source.get("tap_cm")]
        ids, unknown = set(), False
        for value in filter(None, cm_values):
            # Never silently prefer CM over TAP CM, or a first name over the rest.
            for token in re.split(r"[,，;；/\n]+", value):
                matched = aliases.get(token.strip().casefold(), set())
                ids.update(matched)
                unknown |= len(matched) != 1
        responsible = next(iter(ids)) if len(ids) == 1 and not unknown else None
        if not responsible:
            reasons.append("ownership_uncertain" if any(cm_values) else "missing_cm")
        # No guessing another CM for a self-service upload. Preserve as a visible row error.
        forbidden = bool(not scope["is_owner"] and ids and ids != {scope["actor_id"]})
        if forbidden:
            result.append(
                {
                    **source,
                    "handle": handle,
                    "status": "invalid",
                    "reasons": ["文件中这一行属于其他负责人，未导入。"],
                    "cm_name": " / ".join(filter(None, cm_values)),
                }
            )
            continue
        contact_id, contact_name, stable_open_id, trusted = None, None, None, False
        relationship = {
            "本人": "owner",
            "owner": "owner",
            "代表": "representative",
            "representative": "representative",
        }.get(source.get("relationship"))
        stable = source.get("contact_id") or source.get("open_id")
        if responsible and stable and not source.get("formula"):
            try:
                uuid = UUID(source["contact_id"]) if source.get("contact_id") else None
                users = workflow.selectable_users(
                    conn, UUID(responsible), query=source.get("open_id") or "", contact_id=uuid, limit=200
                )
                ids_found = {
                    u["contact_id"] for u in users if not source.get("open_id") or u["open_id"] == source["open_id"]
                }
                if len(ids_found) == 1:
                    contact_id = next(iter(ids_found))
                    contact_name = next(u["display_name"] for u in users if u["contact_id"] == contact_id)
                    stable_open_id = next(u["open_id"] for u in users if u["contact_id"] == contact_id)
                    trusted = True
            except (ValueError, TypeError):
                pass
        if not stable and handle and responsible:
            matches = conn.execute(
                """SELECT link.contact_id,link.relationship,contact.display_name
                FROM social_accounts account
                JOIN contact_social_accounts link ON link.social_account_id=account.id
                JOIN contacts contact ON contact.id=link.contact_id
                WHERE account.platform='tiktok' AND account.normalized_handle=%s
                  AND link.status='confirmed' AND link.relationship='owner' """,
                (handle,),
            ).fetchall()
            if len(matches) == 1 and workflow.selectable_users(
                conn, UUID(responsible), contact_id=matches[0]["contact_id"]
            ):
                contact_id, contact_name, trusted = str(matches[0]["contact_id"]), matches[0]["display_name"], True
                if relationship and relationship != matches[0]["relationship"]:
                    reasons.append("identity_uncertain")
                    trusted = False
                else:
                    relationship = matches[0]["relationship"]
        if not contact_id:
            reasons.append("identity_uncertain" if stable else "missing_user")
        if not relationship:
            reasons.append("missing_relationship")
        if (
            contact_id
            and conn.execute(
                """SELECT 1 FROM audit_events
            WHERE action='creator_handle_identity_review_flagged' AND metadata->>'contact_id'=%s
            AND (metadata->>'handle'=%s OR metadata->>'handle' IS NULL) LIMIT 1""",
                (contact_id, handle),
            ).fetchone()
        ):
            reasons.append("identity_risk")
        if source.get("formula"):
            handle, contact_id, trusted = None, None, False
            reasons.extend(["handle_uncertain", "identity_uncertain"])
        if handle:
            existing = conn.execute(
                """SELECT responsible_authorization_id,contact_id FROM creator_account_assignments a
                JOIN social_accounts s ON s.id=a.social_account_id
                WHERE s.platform='tiktok' AND s.normalized_handle=%s AND a.ended_at IS NULL""",
                (handle,),
            ).fetchone()
            if existing and str(existing["responsible_authorization_id"]) != responsible:
                reasons.append("ownership_conflict")
            if existing and contact_id and str(existing["contact_id"]) != contact_id:
                reasons.append("identity_conflict")
        reasons = list(dict.fromkeys(reasons))
        result.append(
            {
                **source,
                "original_handle": source.get("handle"),
                "handle": handle,
                "responsible_id": responsible,
                "contact_id": contact_id,
                "contact_name": contact_name,
                "cm_name": " / ".join(filter(None, cm_values)) or None,
                "stable_open_id": stable_open_id,
                "relationship": relationship,
                "trusted_identity": trusted,
                "status": "ready" if not reasons else "pending",
                "reasons": reasons,
            }
        )
    return result


def preview(scope, content: bytes, filename: str):
    filename = PurePath(filename.replace("\\", "/")).name[:200]
    batch_id = uuid4()
    with connection() as conn, conn.transaction():
        current = workflow._fresh(conn, scope)
        if not current["can_import"]:
            raise PermissionError("当前页面为只读，不能导入。")
        rows = parse_workbook(content, filename)
        resolved = resolve_rows(conn, current, rows)
        conn.execute(
            """INSERT INTO creator_workflow_import_batches(id,actor_id,responsible_id,filename,file_hash,rows)
            VALUES(%s,%s,%s,%s,%s,%s::jsonb)""",
            (
                batch_id,
                current["actor_id"],
                current["responsible_id"],
                filename,
                hashlib.sha256(content).hexdigest(),
                json.dumps(rows, ensure_ascii=False),
            ),
        )
    return {
        "batch_id": str(batch_id),
        "filename": filename,
        "rows": resolved,
        "summary": {
            **{key: sum(r["status"] == key for r in resolved) for key in ("ready", "pending", "invalid")},
            "total": len(rows),
        },
    }


def commit(scope, batch_id: UUID, request_id: UUID):
    workflow._uuid(request_id)
    with connection() as conn, conn.transaction():
        current = workflow._fresh(conn, scope)
        if not current["can_import"]:
            raise PermissionError("当前页面为只读，不能导入。")
        batch = conn.execute(
            """SELECT *,expires_at>NOW() AS valid FROM creator_workflow_import_batches
            WHERE id=%s AND actor_id=%s AND responsible_id IS NOT DISTINCT FROM %s::uuid FOR UPDATE""",
            (batch_id, current["actor_id"], current["responsible_id"]),
        ).fetchone()
        if not batch:
            raise LookupError("找不到本人的导入预检，请重新上传。")
        if batch["result"] is not None:
            return batch["result"]
        if not batch["valid"]:
            raise workflow.WorkflowInputError("导入预检已过期，请重新上传。")
        resolved = resolve_rows(conn, current, batch["rows"])
        result = {"imported": 0, "formal": 0, "pending": 0, "skipped": 0, "items": []}
        for row in resolved:
            item = {"sheet": row["sheet"], "row": row["row"]}
            if row["status"] == "invalid":
                result["skipped"] += 1
                result["items"].append({**item, "error": row["reasons"][0]})
                continue
            try:
                with conn.transaction():
                    record = workflow.ingest_source(
                        UUID(current["actor_id"]),
                        source_type="excel",
                        source_key=f"{batch['file_hash']}:{row['sheet']}:{row['row']}",
                        source_data={
                            "filename": batch["filename"],
                            "sheet": row["sheet"],
                            "row": row["row"],
                            "file_hash": batch["file_hash"],
                            "cm": row.get("cm"),
                            "tap_cm": row.get("tap_cm"),
                            "original_handle": row.get("original_handle"),
                            "stable_open_id": row.get("stable_open_id"),
                        },
                        responsible_id=row.get("responsible_id"),
                        handle=row.get("handle"),
                        contact_id=row.get("contact_id"),
                        uncertainties=row["reasons"],
                        relationship=row.get("relationship"),
                        trusted_identity=row["trusted_identity"],
                        conn=conn,
                    )
                result["imported"] += 1
                result["formal" if record["status"] == "confirmed" else "pending"] += 1
                result["items"].append({**item, "id": record["id"], "status": record["status"]})
            except (workflow.WorkflowInputError, PermissionError, LookupError) as exc:
                result["skipped"] += 1
                result["items"].append({**item, "error": str(exc)})
        conn.execute(
            "UPDATE creator_workflow_import_batches SET result=%s::jsonb WHERE id=%s",
            (json.dumps(result, default=str, ensure_ascii=False), batch_id),
        )
        return result
