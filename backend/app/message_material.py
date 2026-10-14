"""Pure, bounded extraction of untrusted message bodies and attachment bytes."""

import base64
import csv
import io
import json
import warnings
from datetime import UTC, datetime
from pathlib import PurePosixPath
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from PIL import Image, UnidentifiedImageError

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_TEXT_CHARS = 40_000
MAX_ROWS = 1000
MAX_IMAGE_PIXELS = 20_000_000
_LOCALES = ("zh_cn", "en_us", "en_gb", "ja_jp", "zh_hk", "zh_tw")
_WORD = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


class _Unreadable(ValueError):
    pass


def _parsed(raw):
    if isinstance(raw, dict | list):
        return raw
    if not isinstance(raw, str):
        return None
    try:
        return json.loads(raw)
    except (ValueError, RecursionError):
        return raw


def _language(value):
    """Locale variants are alternate renderings, not separate messages."""
    if isinstance(value, dict):
        for locale in _LOCALES:
            if value.get(locale):
                return value[locale]
        for key, child in value.items():
            if len(key) == 5 and key[2] == "_" and isinstance(child, dict | list):
                return child
    return value


def _texts(value):
    value = _language(value)
    if isinstance(value, list):
        return [text for child in value for text in _texts(child)]
    if not isinstance(value, dict):
        return []
    tag = value.get("tag")
    if tag in ("plain_text", "lark_md", "markdown"):
        content = value.get("content")
        return [content] if isinstance(content, str) and content else []
    if tag == "a":
        text, href = value.get("text"), value.get("href")
        if isinstance(href, str) and href:
            return [f"{text} ({href})" if text and text != href else href]
        return [text] if isinstance(text, str) and text else []
    result = []
    title = value.get("title")
    if isinstance(title, str) and title and isinstance(value.get("content"), list):
        result.append(title)
    text = value.get("text")
    if isinstance(text, str) and text:
        result.append(text)
    for key, child in value.items():
        if key in ("value", "behaviors", "config", "extra"):
            continue
        if key == "i18n_elements":
            result.extend(_texts(_language(child)))
        elif key == "elements" and value.get("i18n_elements"):
            continue
        elif key not in ("text", "title") or not isinstance(child, str):
            result.extend(_texts(child))
    return result


def _nodes(value):
    if isinstance(value, dict):
        yield value
        for child in value.values():
            yield from _nodes(child)
    elif isinstance(value, list):
        for child in value:
            yield from _nodes(child)


def _limited(text, notes):
    lines = text.splitlines()
    if len(lines) > MAX_ROWS:
        text = "\n".join(lines[:MAX_ROWS])
        notes.append("内容超过 1000 行，仅提取前 1000 行")
    if len(text) > MAX_TEXT_CHARS:
        text = text[:MAX_TEXT_CHARS]
        notes.append("文字过长，仅提取前 40000 字符")
    return text


def text_material(raw, message_type):
    """Return locally readable text and explicit gaps; never follow links or IDs."""
    value, notes = _parsed(raw), []
    kind = message_type or "unknown"
    if kind == "merge_forward":
        items = next((value[key] for key in ("items", "messages")
                      if isinstance(value, dict) and isinstance(value.get(key), list)), None)
        text, notes = _forward_material(items, value.get("message_id")) if items else ("", [])
        if not text and not notes:
            notes.append("合并转发内容尚未展开")
    elif kind in ("audio", "media", "video", "sticker"):
        text = ""
        notes.append("音频暂未转写" if kind == "audio" else
                     "表情包暂不支持读取" if kind == "sticker" else "视频暂未解析")
    else:
        text = value if isinstance(value, str) and kind == "text" else "\n".join(_texts(value))
        if kind == "interactive" and any(
            node.get("tag") in ("img", "image", "file", "media") for node in _nodes(value)
        ):
            notes.append("卡片中的图片或附件暂未读取")
        if kind == "post" and any(node.get("tag") in ("audio", "media", "video", "sticker")
                                  for node in _nodes(value)):
            notes.append("富文本中的音视频或表情包暂未解析")
        if not text and kind not in ("image", "file"):
            notes.append("消息正文暂无法解析")
    return {"text": _limited(text, notes), "notes": list(dict.fromkeys(notes))}


def render_forward(items):
    """Expanded message records are quotes; their sender is not the forwarder."""
    return _forward_material(items)[0]


def _forward_material(items, container_id=None):
    parts, notes = [], []
    expanded_ids = {item.get("upper_message_id") for item in items or [] if isinstance(item, dict)
                    and item.get("upper_message_id")}
    for item in items or []:
        if not isinstance(item, dict):
            continue
        body = item.get("body") or {}
        raw = body.get("content", "") if isinstance(body, dict) else body
        kind = item.get("msg_type") or item.get("message_type") or "text"
        if kind == "merge_forward":
            if (container_id and item.get("message_id") == container_id) or item.get("message_id") in expanded_ids:
                continue
        material = text_material(raw, kind)
        if not material["text"] and not material["notes"]:
            material["notes"].append("该引用含未解析附件")
        notes.extend(material["notes"])
        sender = item.get("sender") or {}
        name = ((sender.get("name") or sender.get("sender_name") or sender.get("id"))
                if isinstance(sender, dict) else None)
        name = name or item.get("sender_name") or "未知发送者"
        stamp = item.get("create_time") or item.get("sent_at") or "未知时间"
        if str(stamp).isdigit():
            try:
                stamp = datetime.fromtimestamp(int(stamp) / 1000, UTC).isoformat(timespec="seconds")
            except (ValueError, OverflowError, OSError):
                stamp = "未知时间"
        details = material["text"]
        if material["notes"]:
            details += "\n[读取缺口：" + "；".join(material["notes"]) + "]"
        parts.append(f"引用发送者：{name} · {stamp}\n{details.strip()}")
    text = "以下为转发引用，不是转发者本人陈述：\n" + "\n\n".join(parts) if parts else ""
    return text, list(dict.fromkeys(notes))


def resource_refs(raw, message_type):
    """Only supported IM resources; no audio/video/sticker/card downloads."""
    value = _parsed(raw)
    if not isinstance(value, dict) or message_type not in ("image", "file", "post"):
        return []
    refs, seen = [], set()
    nodes = _nodes(value) if message_type == "post" else (value,)
    for node in nodes:
        kind = message_type if message_type != "post" else node.get("tag")
        if node.get("is_folder") or kind in ("audio", "media", "video", "sticker"):
            continue
        key = node.get("image_key") if kind in ("image", "img") else node.get("file_key")
        resource_type = "image" if kind in ("image", "img") else "file"
        if resource_type == "file" and message_type == "post" and kind not in (None, "file"):
            continue
        if isinstance(key, str) and key and (resource_type, key) not in seen:
            seen.add((resource_type, key))
            name = node.get("file_name") or node.get("name") or ""
            refs.append({"key": key, "type": resource_type, "name": name if isinstance(name, str) else ""})
    return refs


def _archive(data):
    archive = ZipFile(io.BytesIO(data))
    entries = archive.infolist()
    if len(entries) > 10_000 or sum(item.file_size for item in entries) > 50 * 1024 * 1024:
        archive.close()
        raise _Unreadable("压缩后的文档展开过大，未解析")
    if any(item.flag_bits & 1 or item.file_size > 20 * 1024 * 1024 for item in entries):
        archive.close()
        raise _Unreadable("文档含加密或过大的内容，未解析")
    return archive


def _check_xml(data):
    markup = data.replace(b"\x00", b"")  # Also reject declarations in UTF-16/UTF-32 XML.
    if b"<!DOCTYPE" in markup or b"<!ENTITY" in markup:
        raise _Unreadable("文档含不支持的 XML 实体，未解析")


def _xml(data):
    _check_xml(data)
    return ElementTree.fromstring(data)


def _docx_text(data, notes):
    with _archive(data) as archive:
        root = _xml(archive.read("word/document.xml"))
        names = archive.namelist()
        if any(name.startswith("word/media/") for name in names):
            notes.append("文档内嵌图片暂未解析")
        if any(name.startswith(("word/header", "word/footer", "word/footnotes", "word/endnotes", "word/comments"))
               for name in names):
            notes.append("仅提取文档正文，页眉页脚、脚注或批注未读取")

        def inline(node):
            if node.tag in (_WORD + "drawing", _WORD + "object", _WORD + "pict"):
                return ""
            if node.tag == _WORD + "t":
                return node.text or ""
            if node.tag in (_WORD + "br", _WORD + "cr"):
                return "\n"
            if node.tag == _WORD + "tab":
                return "\t"
            return "".join(inline(child) for child in node)

        return "\n".join(inline(node) for node in root.iter(_WORD + "p"))


def _xlsx_text(data, notes):
    with _archive(data) as archive:
        names = archive.namelist()
        if "xl/workbook.xml" not in names or any(name.lower().endswith("vbaproject.bin") for name in names):
            raise _Unreadable("不支持该工作簿或含宏文件")
        if any(name.startswith("xl/media/") for name in names):
            notes.append("工作簿内嵌图片暂未解析")
        # openpyxl parses XML too; reject entities before passing untrusted parts to it.
        for name in names:
            if name.endswith(".xml"):
                _check_xml(archive.read(name))
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    lines, rows, formula = [], 0, False
    try:
        for sheet in workbook.worksheets:
            lines.append(f"工作表：{sheet.title}")
            remaining = MAX_ROWS - rows
            if remaining <= 0:
                notes.append("工作簿超过 1000 行，仅提取前 1000 行")
                break
            if sheet.max_column and sheet.max_column > 200:
                notes.append(f"工作表「{sheet.title}」仅提取前 200 列")
            count = 0
            for row in sheet.iter_rows(max_row=min(sheet.max_row or remaining + 1, remaining + 1),
                                       max_col=min(sheet.max_column or 200, 200), values_only=True):
                if count == remaining:
                    if any(value is not None for value in row) or (sheet.max_row or 0) > remaining:
                        notes.append("工作簿超过 1000 行，仅提取前 1000 行")
                    break
                values = [str(value) if value is not None else "" for value in row]
                formula |= any(value.startswith("=") for value in values)
                lines.append("\t".join(values).rstrip())
                count += 1
            rows += count
    finally:
        workbook.close()
    if formula:
        notes.append("工作簿公式保留原式，未执行计算")
    return "\n".join(lines)


def _decode_text(data):
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    if b"\x00" in data:
        raise _Unreadable("文件不是可读取的纯文本")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("gb18030")


def file_inputs(data, filename, mime):
    """Produce inline model inputs or local text; never write/upload/execute files."""
    result = {"text": "", "inputs": [], "notes": [], "supported": False}
    notes = result["notes"]
    name = PurePosixPath((filename or "attachment").replace("\\", "/")).name[:200] or "attachment"
    suffix = PurePosixPath(name).suffix.lower()
    mime = (mime or "").split(";", 1)[0].lower()
    if not isinstance(data, bytes) or not data:
        notes.append("附件为空，未解析")
        return result
    if len(data) > MAX_FILE_BYTES:
        notes.append("附件超过 20 MB，未解析")
        return result
    try:
        if mime.startswith("image/") or suffix in (".jpg", ".jpeg", ".png", ".webp", ".gif", ".bmp", ".tif", ".tiff"):
            with warnings.catch_warnings():
                warnings.simplefilter("error", Image.DecompressionBombWarning)
                with Image.open(io.BytesIO(data)) as image:
                    if image.width * image.height > MAX_IMAGE_PIXELS or image.width <= 0 or image.height <= 0:
                        raise _Unreadable("图片像素过大，未解析")
                    image_format = image.format
                    if image_format not in ("JPEG", "PNG", "WEBP", "GIF", "BMP", "TIFF"):
                        raise _Unreadable("暂不支持该图片格式")
                    image.verify()
                with Image.open(io.BytesIO(data)) as image:
                    if image_format in ("GIF", "BMP", "TIFF") or getattr(image, "is_animated", False):
                        if getattr(image, "is_animated", False):
                            notes.append("动态图仅解析第一帧")
                        output = io.BytesIO()
                        image.convert("RGBA").save(output, format="PNG")
                        data, image_format = output.getvalue(), "PNG"
                    else:
                        image.load()
                if len(data) > MAX_FILE_BYTES:
                    raise _Unreadable("转换后的图片过大，未解析")
            image_mime = {"JPEG": "image/jpeg", "PNG": "image/png", "WEBP": "image/webp"}[image_format]
            result["inputs"] = [{"type": "input_image", "detail": "auto", "image_url":
                                 f"data:{image_mime};base64,{base64.b64encode(data).decode('ascii')}"}]
        elif suffix == ".pdf" or mime == "application/pdf":
            if not data.lstrip().startswith(b"%PDF-"):
                raise _Unreadable("PDF 文件格式无效，未解析")
            result["inputs"] = [{"type": "input_file", "filename": name if suffix == ".pdf" else name + ".pdf",
                                 "file_data": "data:application/pdf;base64," + base64.b64encode(data).decode("ascii")}]
        elif suffix == ".docx":
            result["text"] = _docx_text(data, notes)
        elif suffix == ".xlsx":
            result["text"] = _xlsx_text(data, notes)
        elif suffix in (".csv", ".txt", ".md") or mime in ("text/plain", "text/csv", "text/markdown"):
            text = _decode_text(data)
            if suffix == ".csv" or mime == "text/csv":
                rows = []
                for index, row in enumerate(csv.reader(io.StringIO(text))):
                    if index >= MAX_ROWS:
                        notes.append("CSV 超过 1000 行，仅提取前 1000 行")
                        break
                    rows.append("\t".join(row))
                text = "\n".join(rows)
            result["text"] = text
        else:
            notes.append("暂不支持该附件格式")
            return result
        result["text"] = _limited(result["text"], notes)
        result["supported"] = bool(result["inputs"] or result["text"].strip())
        if not result["supported"]:
            notes.append("附件没有可提取的文字")
    except _Unreadable as exc:
        result["text"], result["inputs"] = "", []
        notes.append(str(exc))
    except (ValueError, TypeError, KeyError, OSError, RuntimeError, RecursionError, BadZipFile,
            ElementTree.ParseError, csv.Error, UnidentifiedImageError, Image.DecompressionBombError,
            Image.DecompressionBombWarning):
        result["text"], result["inputs"] = "", []
        notes.append("附件损坏、加密或格式不受支持，未解析")
    result["notes"] = list(dict.fromkeys(notes))
    return result
